"""Audit a .pptx for the technical-pptx-design checks that can be automated.

Usage: python audit_pptx.py deck.pptx [--font Calibri] [--code-font Consolas] [--min-pt 11] [--min-dpi 200]
Reports, per slide: shapes outside the slide, overlapping shapes, text below the minimum size (as displayed: inherited
sizes and autofit shrinking included), text set in a font other than the deck's family (or the equation font, or the
code font if given), raster images below the minimum effective resolution (cropping included), empty placeholders,
animations and transitions; and, for the deck: the theme fonts, fonts in use, slide layouts no slide uses, and how many
charts, tables, and vector pictures are native. Text overflow and the text inside charts are not checked here: look at
the rendered slides.
"""
import argparse
import collections
import itertools

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.shapes.picture import Picture
from pptx.util import Emu

EMU_PER_INCH = 914400
DEFAULT_PT = 18  # PowerPoint's size for text no style sets
NS_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
NS_A = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
SVG_BLIP = "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip"
THEME_REFERENCES = {"+mj-lt", "+mn-lt"}  # "use the theme's heading or body font"
RASTER_TYPES = {"image/png", "image/jpeg", "image/tiff", "image/bmp", "image/gif"}
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_BODY,
              PP_PLACEHOLDER.VERTICAL_OBJECT}


def theme_fonts(deck):
    """The Latin heading (major) and body (minor) fonts of each slide master's theme."""
    fonts = []
    for master in deck.slide_masters:
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        fonts.append((theme.find(".//a:majorFont/a:latin", NS_A).get("typeface"),
                      theme.find(".//a:minorFont/a:latin", NS_A).get("typeface")))
    return fonts


def box(shape):
    return shape.left or 0, shape.top or 0, (shape.left or 0) + (shape.width or 0), (shape.top or 0) + (shape.height or 0)


def overlap_area(a, b):
    width = min(a[2], b[2]) - max(a[0], b[0])
    height = min(a[3], b[3]) - max(a[1], b[1])
    return width * height if width > 0 and height > 0 else 0


def text_bodies(shape):
    """The p:txBody / a:txBody elements of a shape: its own text, or each table cell's."""
    if shape.has_text_frame:
        yield shape.text_frame._txBody
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame._txBody


def list_styles(shape, deck):
    """The list styles a shape's text inherits from, nearest first (excluding the text body's own)."""
    if shape.is_placeholder:
        layout_placeholder = getattr(shape, "_base_placeholder", None)
        master_placeholder = getattr(layout_placeholder, "_base_placeholder", None)
        for placeholder in (layout_placeholder, master_placeholder):
            if placeholder is not None:
                yield placeholder._element.find(f"{NS_P}txBody/a:lstStyle", NS_A)
        kind = shape.placeholder_format.type
        style = "titleStyle" if kind in TITLE_TYPES else "bodyStyle" if kind in BODY_TYPES else "otherStyle"
        yield shape.part.slide_layout.slide_master._element.find(f"{NS_P}txStyles/{NS_P}{style}")
    else:
        yield deck.part._element.find(f"{NS_P}defaultTextStyle")


def level_size(list_style, level):
    """The size (pt) a list style sets for a paragraph level, or None."""
    if list_style is None:
        return None
    properties = list_style.find(f"a:lvl{level + 1}pPr/a:defRPr", NS_A)
    return int(properties.get("sz")) / 100 if properties is not None and properties.get("sz") else None


def displayed_sizes(shape, deck):
    """Yield (size in pt as displayed, autofit scale, run) for each text run of the shape."""
    inherited = list(list_styles(shape, deck))
    for body in text_bodies(shape):
        autofit = body.find("a:bodyPr/a:normAutofit", NS_A)
        scale = int(autofit.get("fontScale", "100000")) / 100000 if autofit is not None else 1
        own_style = body.find("a:lstStyle", NS_A)
        for p in body.findall("a:p", NS_A):
            properties = p.find("a:pPr", NS_A)
            level = int(properties.get("lvl", "0")) if properties is not None else 0
            styled = next((size for size in (level_size(style, level) for style in [own_style, *inherited]) if size),
                          DEFAULT_PT)
            for r in p.findall("a:r", NS_A):
                run_properties = r.find("a:rPr", NS_A)
                size = int(run_properties.get("sz")) / 100 if run_properties is not None and run_properties.get("sz") \
                    else styled
                yield size * scale, scale, r


def run_font(r):
    latin = r.find("a:rPr/a:latin", NS_A)
    return latin.get("typeface") if latin is not None else None


def run_text(r):
    return "".join(r.itertext())[:40]


def audit(path, font="Calibri", math_font="Cambria Math", code_font=None, min_pt=11, min_dpi=200):
    """Return {"slides": {number: [notes]}, "deck": [notes], "summary": {...}}; every note is a problem."""
    deck = Presentation(path)
    allowed_fonts = {font, math_font, code_font} - {None} | THEME_REFERENCES
    slide_w, slide_h = deck.slide_width, deck.slide_height
    fonts = collections.Counter()
    counts = collections.Counter()
    slides = {}

    for number, slide in enumerate(deck.slides, start=1):
        notes = []
        shapes = list(slide.shapes)
        for shape in shapes:
            left, top, right, bottom = box(shape)
            if left < 0 or top < 0 or right > slide_w or bottom > slide_h:
                notes.append(f"outside the slide: {shape.name}")
            smallest = None
            for size, scale, r in displayed_sizes(shape, deck):
                name = run_font(r)
                if name:
                    fonts[name] += 1
                    if name not in allowed_fonts:
                        notes.append(f"font {name!r} instead of {font!r} in {shape.name}: {run_text(r)!r}")
                if size < min_pt and (smallest is None or size < smallest[0]):
                    smallest = size, scale, r
            if smallest:
                size, scale, r = smallest
                shrunk = f" (autofit shrinks it to {scale:.0%})" if scale < 1 else ""
                notes.append(f"text at {size:.3g} pt{shrunk} in {shape.name}: {run_text(r)!r}")
            if isinstance(shape, Picture):  # includes pictures in placeholders
                if shape._element.find(f".//{SVG_BLIP}") is not None:
                    counts["vector pictures"] += 1
                else:
                    counts["raster pictures"] += 1
                    try:
                        image = shape.image
                    except (KeyError, ValueError):  # a linked, not embedded, picture
                        image = None
                    if image is not None and image.content_type in RASTER_TYPES and shape.width and shape.height:
                        px_w, px_h = image.size
                        shown_w = px_w * (1 - shape.crop_left - shape.crop_right)
                        shown_h = px_h * (1 - shape.crop_top - shape.crop_bottom)
                        dpi = min(shown_w / (shape.width / EMU_PER_INCH), shown_h / (shape.height / EMU_PER_INCH))
                        if dpi < min_dpi:
                            notes.append(f"image at {dpi:.0f} dpi as displayed: {shape.name}")
            if getattr(shape, "has_chart", False) and shape.has_chart:
                counts["native charts"] += 1
            if getattr(shape, "has_table", False) and shape.has_table:
                counts["native tables"] += 1
            if shape.is_placeholder and shape.has_text_frame and not shape.text_frame.text.strip():
                notes.append(f"empty placeholder: {shape.name}")
        for a, b in itertools.combinations(shapes, 2):
            area = overlap_area(box(a), box(b))
            smaller = min((a.width or 0) * (a.height or 0), (b.width or 0) * (b.height or 0)) or 1
            if area / smaller > 0.02:
                notes.append(f"overlap {area / smaller:.0%} of the smaller: {a.name} / {b.name}")
        xml = slide._element
        # PowerPoint wraps newer animations and transitions in mc:AlternateContent, so search the whole slide
        if xml.find(f".//{NS_P}timing") is not None:
            notes.append("has animations (p:timing)")
        if xml.find(f".//{NS_P}transition") is not None:
            notes.append("has a transition")
        if notes:
            slides[number] = notes

    deck_notes = []
    themes = theme_fonts(deck)
    for heading, body in themes:
        for role, name in (("heading", heading), ("body", body)):
            if name != font:
                deck_notes.append(f"the theme's {role} font is {name!r}, not {font!r}: set it, so all text inherits "
                                  "the family")
    summary = {
        "slide size": f"{Emu(slide_w).inches:.3f} x {Emu(slide_h).inches:.3f} in",
        "theme fonts": themes,
        "fonts set on runs": dict(fonts),
        "counts": dict(counts),
        "layouts no slide uses": [layout.name for master in deck.slide_masters for layout in master.slide_layouts
                                  if not layout.used_by_slides],
    }
    return {"slides": slides, "deck": deck_notes, "summary": summary}


def problem_count(report):
    return sum(map(len, report["slides"].values())) + len(report["deck"])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("deck")
    parser.add_argument("--font", default="Calibri", help="the deck's font family (default: Calibri)")
    parser.add_argument("--math-font", default="Cambria Math", help="the font PowerPoint equations use")
    parser.add_argument("--code-font", help="the one monospace family for code, if the deck shows code")
    parser.add_argument("--min-pt", type=float, default=11)
    parser.add_argument("--min-dpi", type=float, default=200)
    args = parser.parse_args()

    report = audit(args.deck, args.font, args.math_font, args.code_font, args.min_pt, args.min_dpi)
    for number, notes in report["slides"].items():
        print(f"slide {number}:")
        for note in notes:
            print(f"  - {note}")
    summary = report["summary"]
    print("slide size:", summary["slide size"])
    for heading, body in summary["theme fonts"]:
        print(f"theme fonts: heading {heading!r}, body {body!r}")
    for note in report["deck"]:
        print(f"  - {note}")
    print("fonts set on runs:", summary["fonts set on runs"] or "none (inherited from the theme)")
    print("counts:", summary["counts"])
    print("layouts no slide uses:", summary["layouts no slide uses"] or "none")
    print(f"{problem_count(report)} problem(s)")


if __name__ == "__main__":
    main()
