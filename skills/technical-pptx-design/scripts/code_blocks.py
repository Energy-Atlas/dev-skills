"""Code blocks for the technical-pptx-design skill: listings and comparisons as native, editable PowerPoint tables.

Needs python-pptx and Pygments. Import from this folder:

    import sys
    sys.path.insert(0, "<path to this skill>/scripts")
    from code_blocks import add_code_block, add_code_diff, monospace_font

- monospace_font(): the first installed family of MONOSPACE_FAMILIES (Cascadia Mono, Consolas, ...).
- add_code_block(slide, code, language, ...): one listing, with line numbers, syntax highlighting, a grey backdrop, an
  optional label row, and optional green or red line marks.
- add_code_diff(slide, old, new, language, ...): two versions of the same code compared git-diff style, in one unified
  block or as two aligned blocks side by side; removed lines red with "-", added lines green with "+".

Each block is a table with one row per line: a fill can then mark a whole line, and everything stays editable. The code is
shown exactly as given, except that tabs are expanded to four spaces.
"""
import difflib
import shutil
import subprocess
import sys
import warnings
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pygments.lexers import get_lexer_by_name
from pygments.token import Comment, Keyword, Name, Number, String

MONOSPACE_FAMILIES = ("Cascadia Mono", "Consolas", "JetBrains Mono", "SF Mono", "Menlo", "DejaVu Sans Mono",
                      "Liberation Mono", "Courier New")
NO_STYLE_NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # PowerPoint's table style without fills or borders
CHAR_WIDTH = 0.6  # width of a monospace character, in ems
LINE_HEIGHT = 1.35  # row height, in ems

# GitHub's Primer colours: every text colour reaches 4.5:1 on every backdrop of its theme, line marks included. Syntax
# colours avoid green and red, which mean added and removed.
LIGHT = {
    "backdrop": "F6F8FA", "label": "EAEEF2", "green": "DAFBE1", "red": "FFEBE9",
    "text": "1F2328", "muted": "59636E", "comment": "59636E", "keyword": "6639BA", "string": "0A3069",
    "number": "0550AE", "function": "953800", "added": "116329", "removed": "82071E",
}
DARK = {
    "backdrop": "161B22", "label": "21262D", "green": "12261E", "red": "25171C",
    "text": "E6EDF3", "muted": "9198A1", "comment": "9198A1", "keyword": "D2A8FF", "string": "A5D6FF",
    "number": "79C0FF", "function": "FFA657", "added": "56D364", "removed": "FF7B72",
}


def installed_families():
    """The font families installed on this machine, as well as they can be found."""
    if sys.platform == "win32":
        import winreg
        names = set()
        for root in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
            try:
                key = winreg.OpenKey(root, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts")
            except OSError:
                continue
            index = 0
            while True:
                try:
                    names.add(winreg.EnumValue(key, index)[0].split(" (")[0])
                except OSError:
                    break
                index += 1
        return names
    if shutil.which("fc-list"):
        output = subprocess.run(["fc-list", ":", "family"], capture_output=True, text=True).stdout
        return {family.strip() for line in output.splitlines() for family in line.split(",")}
    folders = [Path("/System/Library/Fonts"), Path("/Library/Fonts"), Path.home() / "Library" / "Fonts"]
    return {path.stem for folder in folders if folder.is_dir() for path in folder.iterdir()}


def monospace_font(candidates=MONOSPACE_FAMILIES, installed=None):
    """The first candidate family that is installed (a registry or file name starting with it counts)."""
    installed = installed_families() if installed is None else installed
    for family in candidates:
        if any(name == family or name.startswith(family + " ") for name in installed):
            return family
    return candidates[-1]


def _category(token_type):
    if token_type in Comment:
        return "comment"
    if token_type in Keyword.Constant or token_type in Number or token_type in Name.Constant \
            or token_type in Name.Builtin:
        return "number"
    if token_type in Keyword or token_type in Name.Namespace:
        return "keyword"
    if token_type in String:
        return "string"
    if token_type in Name.Function or token_type in Name.Class or token_type in Name.Decorator:
        return "function"
    return "text"


def highlighted_lines(code, language):
    """Split code into lines of (category, text) pieces; lexing the whole text keeps multi-line strings right."""
    code = code.expandtabs(4)
    lexer = get_lexer_by_name(language, stripnl=False, stripall=False, ensurenl=False)
    lines = [[]]
    for token_type, value in lexer.get_tokens(code):
        parts = value.split("\n")
        for index, part in enumerate(parts):
            if index:
                lines.append([])
            if part:
                lines[-1].append((_category(token_type), part))
    if code.endswith("\n") and not lines[-1]:
        lines.pop()
    return lines


def _fill(cell, colour):
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(colour)


def _write(cell, pieces, font, size, palette, align=PP_ALIGN.LEFT):
    """Write (category, text) pieces into a cell as runs; an empty line keeps its height with one space."""
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Emu(0)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = align
    for category, text in pieces or [("text", " ")]:
        run = paragraph.add_run()
        run.text = text
        run.font.name = font
        run.font.size = size
        run.font.italic = category == "comment"
        run.font.color.rgb = RGBColor.from_string(palette[category])


def _fit(rows, number_columns, width, size, min_size):
    """The largest size up to `size` (not below min_size) at which the longest line fits, and the column widths."""
    digits = max([len(str(n)) for row in rows for n in row["numbers"] if n] or [1])
    longest = max([sum(len(text) for _, text in row["pieces"]) for row in rows] or [1])
    while True:
        char = size * CHAR_WIDTH
        number_width = int(char * (digits + 1) + Inches(0.16))
        code_width = width - number_width * number_columns
        fits = longest * char <= code_width - Inches(0.16)
        if fits or Pt(size.pt - 1) < min_size:
            break
        size = Pt(size.pt - 1)
    if not fits:
        warnings.warn(f"code lines of {longest} characters wrap at {size.pt:g} pt in a {Emu(width).inches:.2f} in block;"
                      " widen the block, show an excerpt, or split it")
    return size, number_width, code_width


def _table(slide, rows, number_columns, left, top, width, font, size, palette, label, min_size):
    """A code table. Each row: {"numbers": [...], "pieces": [...], "fill": None | "green" | "red" | "label"}."""
    size, number_width, code_width = _fit(rows, number_columns, width, size, min_size)

    count = len(rows) + (1 if label else 0)
    row_height = int(size * LINE_HEIGHT)
    shape = slide.shapes.add_table(count, number_columns + 1, left, top, width, row_height * count)
    table = shape.table
    table.first_row = table.horz_banding = False
    shape._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId")).text = NO_STYLE_NO_GRID
    for index in range(number_columns):
        table.columns[index].width = number_width
    table.columns[number_columns].width = code_width
    for row in table.rows:
        row.height = row_height

    offset = 0
    if label:
        first = table.cell(0, 0)
        first.merge(table.cell(0, number_columns))
        _fill(first, palette["label"])
        _write(first, [("muted", label)], font, size, palette)
        offset = 1
    for index, row in enumerate(rows, start=offset):
        fill = palette[row["fill"]] if row["fill"] else palette["backdrop"]
        for column in range(number_columns + 1):
            cell = table.cell(index, column)
            _fill(cell, fill)
            if column < number_columns:
                number = row["numbers"][column]
                _write(cell, [("muted", str(number))] if number else None, font, size, palette, PP_ALIGN.RIGHT)
            else:
                _write(cell, row["pieces"], font, size, palette)
    return shape


def add_code_block(slide, code, language, left, top, width, *, font=None, size=Pt(14), first_line=1, marks=None,
                   label=None, dark=False, min_size=Pt(12)):
    """One listing. marks maps source line numbers (counted from first_line) to "green" or "red"."""
    palette = DARK if dark else LIGHT
    marks = marks or {}
    rows = [{"numbers": [number], "pieces": pieces, "fill": marks.get(number)}
            for number, pieces in enumerate(highlighted_lines(code, language), start=first_line)]
    return _table(slide, rows, 1, left, top, width, font or monospace_font(), size, palette, label, min_size)


def diff_hunks(old_lines, new_lines):
    """("equal", [(i, j), ...]) and ("change", [i, ...], [j, ...]) groups of 0-based line indices."""
    matcher = difflib.SequenceMatcher(a=old_lines, b=new_lines, autojunk=False)
    hunks = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            hunks.append(("equal", list(zip(range(i1, i2), range(j1, j2)))))
        else:
            hunks.append(("change", list(range(i1, i2)), list(range(j1, j2))))
    return hunks


def _with_marker(marker, category, pieces):
    return [(category, marker + " ")] + pieces


def add_code_diff(slide, old, new, language, left, top, width, *, layout="unified", labels=(None, None), gap=Inches(0.25),
                  context=None, font=None, size=Pt(14), old_first_line=1, new_first_line=1, dark=False,
                  min_size=Pt(12)):
    """Two versions of the same code, git-diff style.

    layout "unified": one block with old and new line numbers, removed lines then added lines; "split": two blocks side
    by side (old left, new right) whose rows line up, an empty grey row standing in for a line the other side lacks.
    context: keep only this many unchanged lines around each change, a "⋯" row marking each gap. Returns the shape, or
    the two shapes of a split.
    """
    palette = DARK if dark else LIGHT
    font = font or monospace_font()
    old_lines, new_lines = highlighted_lines(old, language), highlighted_lines(new, language)
    plain = lambda lines: ["".join(text for _, text in line) for line in lines]
    hunks = diff_hunks(plain(old_lines), plain(new_lines))
    gap_row = {"pieces": [("muted", "⋯")], "fill": "label"}

    def unchanged(pairs, first, last):
        if context is None or len(pairs) <= 2 * context + 1 or (first and last):
            return pairs, None
        head = [] if first else pairs[:context]
        tail = [] if last else pairs[-context:]
        return head + [None] + tail, True

    unified, left_rows, right_rows = [], [], []
    for position, hunk in enumerate(hunks):
        if hunk[0] == "equal":
            pairs, _ = unchanged(hunk[1], position == 0, position == len(hunks) - 1)
            for pair in pairs:
                if pair is None:
                    unified.append({"numbers": [None, None], **gap_row})
                    left_rows.append({"numbers": [None], **gap_row})
                    right_rows.append({"numbers": [None], **gap_row})
                    continue
                i, j = pair
                unified.append({"numbers": [i + old_first_line, j + new_first_line],
                                "pieces": _with_marker(" ", "text", old_lines[i]), "fill": None})
                left_rows.append({"numbers": [i + old_first_line], "pieces": _with_marker(" ", "text", old_lines[i]),
                                  "fill": None})
                right_rows.append({"numbers": [j + new_first_line], "pieces": _with_marker(" ", "text", new_lines[j]),
                                   "fill": None})
        else:
            _, removed, added = hunk
            for i in removed:
                unified.append({"numbers": [i + old_first_line, None],
                                "pieces": _with_marker("-", "removed", old_lines[i]), "fill": "red"})
            for j in added:
                unified.append({"numbers": [None, j + new_first_line],
                                "pieces": _with_marker("+", "added", new_lines[j]), "fill": "green"})
            for k in range(max(len(removed), len(added))):
                if k < len(removed):
                    i = removed[k]
                    left_rows.append({"numbers": [i + old_first_line],
                                      "pieces": _with_marker("-", "removed", old_lines[i]), "fill": "red"})
                else:
                    left_rows.append({"numbers": [None], "pieces": [], "fill": "label"})
                if k < len(added):
                    j = added[k]
                    right_rows.append({"numbers": [j + new_first_line],
                                       "pieces": _with_marker("+", "added", new_lines[j]), "fill": "green"})
                else:
                    right_rows.append({"numbers": [None], "pieces": [], "fill": "label"})

    if layout == "unified":
        label = " → ".join(text for text in labels if text) or None
        return _table(slide, unified, 2, left, top, width, font, size, palette, label, min_size)
    if layout != "split":
        raise ValueError(f"layout must be 'unified' or 'split', not {layout!r}")
    half = int((width - gap) / 2)
    # both halves at the size the tighter side needs, so the two blocks match
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        size = min(_fit(left_rows, 1, half, size, min_size)[0], _fit(right_rows, 1, half, size, min_size)[0])
    shapes = (_table(slide, left_rows, 1, left, top, half, font, size, palette, labels[0], min_size),
              _table(slide, right_rows, 1, left + half + gap, top, half, font, size, palette, labels[1], min_size))
    return shapes
