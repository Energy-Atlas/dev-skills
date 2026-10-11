"""python-pptx helpers for the technical-pptx-design skill.

Import them from this folder:

    import sys
    sys.path.insert(0, "<path to this skill>/scripts")
    from pptx_helpers import new_widescreen_deck, set_theme_font, add_svg_picture, remove_unused_layouts

- new_widescreen_deck(): a blank 16:9 deck whose master and layouts are laid out for 13.333 in.
- set_theme_font(deck, family): the theme's heading and body font, so all text inherits the family.
- add_svg_picture(slide, svg, png, ...): a vector (SVG) picture with a PNG fallback, as PowerPoint stores one.
- remove_unused_layouts(deck): drop every slide layout no slide uses.
"""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.ns import qn
from pptx.util import Inches

NS_A = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
NS_ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXTENSION_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
WIDE, HIGH = Inches(13.333), Inches(7.5)


def new_widescreen_deck():
    """A blank 16:9 deck (13.333 x 7.5 in).

    python-pptx's Presentation() is 4:3 (10 x 7.5 in), and resizing the slide alone leaves every layout's placeholders
    laid out for 10 in. This also stretches the shapes of the master and of each layout to the new width.
    """
    deck = Presentation()
    factor = WIDE / deck.slide_width
    deck.slide_width, deck.slide_height = WIDE, HIGH
    for master in deck.slide_masters:
        for shapes in [master.shapes] + [layout.shapes for layout in master.slide_layouts]:
            for shape in shapes:
                # only shapes with their own position; the others inherit it from the (already stretched) master
                for xfrm in shape._element.xpath("./p:spPr/a:xfrm | ./p:xfrm"):
                    offset, extent = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
                    offset.set("x", str(round(int(offset.get("x")) * factor)))
                    extent.set("cx", str(round(int(extent.get("cx")) * factor)))
    return deck


def set_theme_font(deck, family):
    """Set the heading and body Latin fonts of every slide master's theme, so all text inherits the family."""
    for master in deck.slide_masters:
        part = master.part.part_related_by(RT.THEME)
        theme = etree.fromstring(part.blob)
        for latin in theme.findall(".//a:majorFont/a:latin", NS_A) + theme.findall(".//a:minorFont/a:latin", NS_A):
            latin.set("typeface", family)
        part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8", standalone=True)


def add_svg_picture(slide, svg_path, png_path, left, top, width=None, height=None):
    """Add a vector picture: the SVG, with a PNG fallback for applications that cannot show SVG.

    python-pptx cannot insert an SVG itself. PowerPoint (Microsoft 365, 2019 and later) shows the SVG and can convert
    it to editable shapes; older versions show the PNG. Export both from the same source (for example matplotlib's
    savefig to .svg and to .png at 200 dpi or more) so they have the same aspect ratio; size comes from the PNG.
    """
    picture = slide.shapes.add_picture(str(png_path), left, top, width, height)
    package = slide.part.package
    svg = Part(package.next_partname("/ppt/media/image%d.svg"), "image/svg+xml", package, Path(svg_path).read_bytes())
    r_id = slide.part.relate_to(svg, RT.IMAGE)
    blip = picture._element.xpath("./p:blipFill/a:blip")[0]
    extensions = blip.find(qn("a:extLst"))
    if extensions is None:
        extensions = etree.SubElement(blip, qn("a:extLst"))
    extension = etree.SubElement(extensions, qn("a:ext"), uri=SVG_EXTENSION_URI)
    svg_blip = etree.SubElement(extension, f"{{{NS_ASVG}}}svgBlip", nsmap={"asvg": NS_ASVG})
    svg_blip.set(qn("r:embed"), r_id)
    return picture


def remove_unused_layouts(deck):
    """Remove every slide layout that no slide uses; return their names."""
    removed = []
    for master in deck.slide_masters:
        for layout in list(master.slide_layouts):
            if not layout.used_by_slides:
                removed.append(layout.name)
                master.slide_layouts.remove(layout)
    return removed
