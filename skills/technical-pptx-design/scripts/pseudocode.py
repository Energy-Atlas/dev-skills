"""Pseudocode blocks for the technical-pptx-design skill: an algorithm in the frame of a paper's algorithm listing.

Needs python-pptx. Import from this folder:

    import sys
    sys.path.insert(0, "<path to this skill>/scripts")
    from pseudocode import add_pseudocode

The block is a native, editable table in the deck's own font (no font is set on its text), without a backdrop: a caption
row ("Algorithm 1: Name") between two thin rules, unnumbered Input/Output rows, numbered lines, one narrow column per
indent level whose left border draws the indent guide, comments right-aligned after a "▷", and a rule at the bottom.

Write the algorithm as plain text, indented by four spaces (or a tab) per level:

    Input: hourly loads $q_t$ of every zone
    Output: the building's peak load $Q_{max}$
    $Q_{max} \\gets 0$
    for each hour $t$ do                      // one pass over the year
        $Q_t \\gets \\sum_z q_{z,t}$
        if $Q_t > Q_{max}$ then
            $Q_{max} \\gets Q_t$
        end if
    end for
    return $Q_{max}$

- Keywords (KEYWORDS: if, then, for, while, return, end, ...) are set in bold.
- Text between dollar signs is maths: letters in italic, `_x` or `_{...}` a subscript, `^x` or `^{...}` a superscript,
  and a few commands (\\gets, \\leq, \\geq, \\neq, \\sum, \\in, \\infty, \\cdot, \\times) become their symbols.
- "// comment" or "▷ comment" at the end of a line becomes a right-aligned comment.
- Leading "Input:", "Output:", "Require:", "Ensure:", "Data:", or "Result:" lines are the unnumbered header.
"""
import re
import warnings

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from code_blocks import DARK, LIGHT, LISTING_MARKS, NO_STYLE_NO_GRID

KEYWORDS = frozenset({"if", "then", "else", "elif", "for", "each", "all", "in", "to", "downto", "do", "while", "repeat",
                      "until", "return", "end", "function", "procedure", "and", "or", "not", "break", "continue"})
HEADERS = ("Input:", "Output:", "Require:", "Ensure:", "Data:", "Result:")
COMMANDS = {r"\gets": "←", r"\leftarrow": "←", r"\leq": "≤", r"\geq": "≥", r"\neq": "≠", r"\sum": "∑", r"\in": "∈",
            r"\infty": "∞", r"\cdot": "·", r"\times": "×"}
INDENT = Inches(0.3)  # width of one indent level
CHAR_WIDTH = 0.5  # average width of a character of a proportional font, in ems
LINE_HEIGHT = 1.35
TCPR_ORDER = ["lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr", "cell3D", "noFill", "solidFill", "gradFill",
              "blipFill", "pattFill", "grpFill", "headers", "extLst"]
PALETTE_EXTRA = {id(LIGHT): {"rule": "1F2328", "guide": "AFB8C1"}, id(DARK): {"rule": "E6EDF3", "guide": "484F58"}}


def parse(text):
    """Split pseudocode into header lines and body lines of (depth, code, comment)."""
    headers, body = [], []
    for raw in text.expandtabs(4).splitlines():
        stripped = raw.strip()
        if not body and stripped.startswith(HEADERS):
            headers.append(stripped)
            continue
        code, comment = stripped, ""
        for mark in ("//", "▷"):
            outside_maths = [i for i in (m.start() for m in re.finditer(re.escape(mark), stripped))
                             if stripped[:i].count("$") % 2 == 0]
            if outside_maths:
                code, comment = stripped[:outside_maths[0]].rstrip(), stripped[outside_maths[0] + len(mark):].strip()
                break
        depth = (len(raw.expandtabs(4)) - len(raw.expandtabs(4).lstrip(" "))) // 4
        body.append((depth, code, comment))
    while body and not body[-1][1] and not body[-1][2]:
        body.pop()
    return headers, body


def pieces(code, keywords=KEYWORDS):
    """Split a line into (text, style) pieces; style is "", "bold", "italic", "sub", or "sup" (subscripts italic)."""
    result = []
    for index, part in enumerate(re.split(r"(\$[^$]*\$)", code)):
        if index % 2:
            result.extend(_maths(part[1:-1]))
            continue
        for word in re.split(r"(\b[A-Za-z]+\b)", part):
            if word:
                result.append((word, "bold" if word.lower() in keywords else ""))
    return result


def _maths(source):
    for command, symbol in sorted(COMMANDS.items(), key=lambda item: -len(item[0])):
        source = source.replace(command, symbol)
    result, index = [], 0
    while index < len(source):
        char = source[index]
        if char in "_^" and index + 1 < len(source):
            if source[index + 1] == "{":
                end = source.find("}", index + 2)
                end = len(source) if end < 0 else end
                script, index = source[index + 2:end], end + 1
            else:
                script, index = source[index + 1], index + 2
            result.append((script, "sub" if char == "_" else "sup"))
            continue
        result.append((char, "italic" if char.isalpha() else ""))
        index += 1
    merged = []
    for text, style in result:  # join neighbours of the same style into one run
        if merged and merged[-1][1] == style:
            merged[-1] = (merged[-1][0] + text, style)
        else:
            merged.append((text, style))
    return merged


def _border(cell, side, colour, width=Pt(0.75)):
    """Draw one edge of a cell (tcPr children must keep the schema's order)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tag = {"left": "lnL", "right": "lnR", "top": "lnT", "bottom": "lnB"}[side]
    old = tc_pr.find(qn(f"a:{tag}"))
    if old is not None:
        tc_pr.remove(old)
    line = etree.Element(qn(f"a:{tag}"), w=str(int(width)), cap="flat", cmpd="sng", algn="ctr")
    fill = etree.SubElement(line, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr"), val=colour)
    etree.SubElement(line, qn("a:prstDash"), val="solid")
    later = [child for child in tc_pr if etree.QName(child).localname in TCPR_ORDER
             and TCPR_ORDER.index(etree.QName(child).localname) > TCPR_ORDER.index(tag)]
    if later:
        later[0].addprevious(line)
    else:
        tc_pr.append(line)


def _cell(cell, runs, size, colour, align=PP_ALIGN.LEFT, fill=None, italic=False, muted=None):
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Emu(0)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        cell.fill.background()
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = align
    for text, style in runs or [(" ", "")]:
        run = paragraph.add_run()
        run.text = text
        run.font.size = size if style not in ("sub", "sup") else Pt(size.pt)  # baseline offset scales it
        run.font.bold = style == "bold"
        run.font.italic = italic or style in ("italic", "sub", "sup")
        run.font.color.rgb = RGBColor.from_string(muted or colour)
        if style in ("sub", "sup"):
            run.font._rPr.set("baseline", "-25000" if style == "sub" else "30000")


def add_pseudocode(slide, text, left, top, width, *, title=None, size=Pt(18), marks=None, keywords=KEYWORDS,
                   first_line=1, dark=False, min_size=Pt(12)):
    """An algorithm block. title: "Algorithm 1: Name" (bold up to the colon). marks: {line number: "yellow" | "blue"}."""
    palette = {**(DARK if dark else LIGHT), **PALETTE_EXTRA[id(DARK if dark else LIGHT)]}
    marks = marks or {}
    wrong = sorted(set(marks.values()) - set(LISTING_MARKS))
    if wrong:
        raise ValueError(f"line marks are {' or '.join(LISTING_MARKS)}, not {', '.join(wrong)}")
    headers, body = parse(text)
    depth = max([d for d, _, _ in body] or [0])
    has_comments = any(comment for _, _, comment in body)
    numbered = [i for i, (_, code, comment) in enumerate(body) if code or comment]
    digits = len(str(first_line + len(numbered)))

    def plain(line):
        return "".join(t for t, _ in pieces(line, keywords))

    while True:  # shrink until the longest line fits, or the minimum size is reached
        em = size
        number_width = int(em * CHAR_WIDTH * (digits + 1) + Inches(0.16))
        comment_chars = max([len(c) + 2 for _, _, c in body if c] or [0])
        comment_width = min(int(em * CHAR_WIDTH * comment_chars + Inches(0.16)), int(width * 0.4)) if has_comments else 0
        code_width = width - number_width - depth * INDENT - comment_width
        need = max([d * INDENT + len(plain(code)) * em * CHAR_WIDTH + Inches(0.16) for d, code, _ in body] or [0])
        fits = need <= depth * INDENT + code_width
        if fits or Pt(size.pt - 1) < min_size:
            break
        size = Pt(size.pt - 1)
    if not fits:
        warnings.warn(f"pseudocode lines wrap at {size.pt:g} pt in a {Emu(width).inches:.2f} in block; widen it or split it")

    columns = 1 + depth + 1 + (1 if has_comments else 0)
    code_column, last = 1 + depth, columns - 1
    rows = (1 if title else 0) + len(headers) + len(body)
    row_height = int(size * LINE_HEIGHT)
    shape = slide.shapes.add_table(rows, columns, left, top, width, row_height * rows)
    table = shape.table
    table.first_row = table.horz_banding = False
    shape._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId")).text = NO_STYLE_NO_GRID
    table.columns[0].width = number_width
    for k in range(depth):
        table.columns[1 + k].width = INDENT
    table.columns[code_column].width = code_width
    if has_comments:
        table.columns[last].width = comment_width
    for row in table.rows:
        row.height = row_height

    # PowerPoint resolves an edge between two cells from both of them, so each line is set on both sides of its edge.
    def rule(row, side):
        for column in range(columns):
            _border(table.cell(row, column), side, palette["rule"])
            neighbour = row - 1 if side == "top" else row + 1
            if 0 <= neighbour < rows:
                _border(table.cell(neighbour, column), "bottom" if side == "top" else "top", palette["rule"])

    def guide(row, column):
        _border(table.cell(row, column), "left", palette["guide"])
        _border(table.cell(row, column - 1), "right", palette["guide"])

    r = 0
    if title:
        name, colon, rest = title.partition(":")
        table.cell(0, 0).merge(table.cell(0, last))
        _cell(table.cell(0, 0), [(name + colon, "bold"), (rest, "")], size, palette["text"])
        rule(0, "top")
        rule(0, "bottom")
        r = 1
    for header in headers:
        label, _, rest = header.partition(":")
        _cell(table.cell(r, 0), None, size, palette["text"])
        table.cell(r, 1).merge(table.cell(r, last))
        _cell(table.cell(r, 1), [(label + ":", "bold")] + pieces(rest, keywords), size, palette["text"])
        r += 1
    if not title:
        rule(0, "top")
    number = first_line
    for line_depth, code, comment in body:
        fill = None
        label = None
        if code or comment:
            fill = palette[marks[number]] if number in marks else None
            label, number = str(number), number + 1
        _cell(table.cell(r, 0), [(label, "")] if label else None, size, palette["text"], PP_ALIGN.RIGHT, fill,
              muted=palette["muted"])
        start = 1 + line_depth
        if start < code_column:  # the code spans the indent columns deeper than its own level
            table.cell(r, start).merge(table.cell(r, code_column))
        for k in range(line_depth):  # an indent guide for every level the line sits inside
            _cell(table.cell(r, 1 + k), None, size, palette["text"], fill=fill)
            guide(r, 1 + k)
        _cell(table.cell(r, start), pieces(code, keywords) if code else None, size, palette["text"], fill=fill)
        if has_comments:
            _cell(table.cell(r, last), [("▷ " + comment, "")] if comment else None, size, palette["text"],
                  PP_ALIGN.RIGHT, fill, italic=True, muted=palette["muted"])
        r += 1
    rule(rows - 1, "bottom")
    return shape
