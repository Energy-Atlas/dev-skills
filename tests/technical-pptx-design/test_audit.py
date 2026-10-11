"""Regression tests for skills/technical-pptx-design/scripts: each defect deck must be caught, a clean deck must pass.

Run from the repository root: python -m pytest tests (needs python-pptx, Pillow, pytest).
"""
import sys
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).parents[2] / "skills" / "technical-pptx-design" / "scripts"))
from audit_pptx import audit, problem_count  # noqa: E402
from pptx_helpers import add_svg_picture, new_widescreen_deck, remove_unused_layouts, set_theme_font  # noqa: E402

SVG = '<svg xmlns="http://www.w3.org/2000/svg" width="400" height="240"><rect width="400" height="240" fill="#0072B2"/></svg>'
TRANSITION = (
    '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
    '<mc:Choice Requires="p14"><p:transition spd="slow" p14:dur="2000"><p14:vortex dir="r"/></p:transition></mc:Choice>'
    '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>'
)


@pytest.fixture
def png(tmp_path):
    def make(width, height, name="figure.png"):
        path = tmp_path / name
        Image.new("RGB", (width, height), "white").save(path)
        return path
    return make


def layout(deck, name):
    return next(item for item in deck.slide_layouts if item.name == name)


def titled_slide(deck, title="A subject title"):
    slide = deck.slides.add_slide(layout(deck, "Title Only"))
    slide.shapes.title.text = title
    return slide


def run_audit(deck, tmp_path, **options):
    set_theme_font(deck, "Calibri")
    remove_unused_layouts(deck)
    path = tmp_path / "deck.pptx"
    deck.save(path)
    return audit(path, **options)


def notes(report):
    return [note for slide_notes in report["slides"].values() for note in slide_notes] + report["deck"]


def test_widescreen_deck_lays_out_its_layouts_for_16_9():
    deck = new_widescreen_deck()
    assert (deck.slide_width, deck.slide_height) == (Inches(13.333), Inches(7.5))
    title = layout(deck, "Title Only").placeholders[0]
    assert title.left + title.width > Inches(12)
    assert title.left + title.width <= deck.slide_width


def test_clean_deck_passes(tmp_path, png):
    deck = new_widescreen_deck()
    cover = deck.slides.add_slide(layout(deck, "Title Slide"))
    cover.shapes.title.text = "Hourly cooling load, case study building"
    cover.placeholders[1].text = "A. Researcher, Energy Atlas, 2026"

    slide = titled_slide(deck)
    svg = tmp_path / "figure.svg"
    svg.write_text(SVG)
    add_svg_picture(slide, svg, png(100, 60), Inches(0.6), Inches(1.6), width=Inches(7.5))  # low-res fallback is fine
    table = slide.shapes.add_table(3, 2, Inches(8.6), Inches(1.6), Inches(4), Inches(1.5)).table
    table.cell(0, 0).text = "Zone"
    code = slide.shapes.add_textbox(Inches(8.6), Inches(4), Inches(4), Inches(0.6)).text_frame
    code.text = "simulate(zones=12)"
    code.paragraphs[0].runs[0].font.name = "Consolas"

    report = run_audit(deck, tmp_path, code_font="Consolas")
    assert problem_count(report) == 0, notes(report)
    assert report["summary"]["layouts no slide uses"] == []
    assert report["summary"]["counts"] == {"vector pictures": 1, "native tables": 1}


def test_code_font_is_flagged_unless_declared(tmp_path):
    deck = new_widescreen_deck()
    box = titled_slide(deck).shapes.add_textbox(Inches(0.6), Inches(2), Inches(4), Inches(0.6)).text_frame
    box.text = "simulate(zones=12)"
    box.paragraphs[0].runs[0].font.name = "Consolas"
    assert any("'Consolas'" in note for note in notes(run_audit(deck, tmp_path)))


def test_low_resolution_picture_in_a_placeholder_is_flagged(tmp_path, png):
    deck = new_widescreen_deck()
    slide = deck.slides.add_slide(layout(deck, "Picture with Caption"))
    slide.shapes.title.text = "Map of the study area"
    placeholder = next(p for p in slide.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.PICTURE)
    placeholder.insert_picture(str(png(800, 600)))
    for p in list(slide.placeholders):
        if p.has_text_frame and not p.text_frame.text:
            p._element.getparent().remove(p._element)
    assert any("dpi as displayed: Picture Placeholder" in note for note in notes(run_audit(deck, tmp_path)))


def test_cropping_counts_against_resolution(tmp_path, png):
    deck = new_widescreen_deck()
    picture = titled_slide(deck).shapes.add_picture(str(png(2000, 1200)), Inches(0.6), Inches(1.6), Inches(5), Inches(3))
    assert problem_count(run_audit(deck, tmp_path)) == 0  # 400 dpi uncropped
    picture.crop_left, picture.crop_right = 0.4, 0.4  # 400 px shown across 5 in: 80 dpi
    assert any("image at 80 dpi" in note for note in notes(run_audit(deck, tmp_path)))


def test_transitions_and_animations_inside_alternate_content_are_flagged(tmp_path):
    deck = new_widescreen_deck()
    titled_slide(deck)._element.append(etree.fromstring(TRANSITION))
    assert "has a transition" in notes(run_audit(deck, tmp_path))


def test_text_shrunk_by_autofit_is_flagged(tmp_path):
    deck = new_widescreen_deck()
    slide = deck.slides.add_slide(layout(deck, "Title and Content"))
    slide.shapes.title.text = "Too much text"
    body = slide.placeholders[1].text_frame
    body.text = "a long paragraph " * 60
    body._txBody.find(qn("a:bodyPr")).append(etree.Element(qn("a:normAutofit"), fontScale="25000"))
    assert any("text at 8 pt (autofit shrinks it to 25%)" in note for note in notes(run_audit(deck, tmp_path)))


def test_inherited_small_text_is_flagged(tmp_path):
    deck = new_widescreen_deck()
    default_size = deck.part._element.find(qn("p:defaultTextStyle")).find(f"{qn('a:lvl1pPr')}/{qn('a:defRPr')}")
    default_size.set("sz", "900")
    titled_slide(deck).shapes.add_textbox(Inches(0.6), Inches(2), Inches(4), Inches(0.6)).text_frame.text = "Source: x"
    assert any("text at 9 pt in TextBox" in note for note in notes(run_audit(deck, tmp_path)))


def test_explicit_small_text_is_flagged(tmp_path):
    deck = new_widescreen_deck()
    box = titled_slide(deck).shapes.add_textbox(Inches(0.6), Inches(2), Inches(4), Inches(0.6)).text_frame
    box.text = "Source: x"
    box.paragraphs[0].runs[0].font.size = Pt(10)
    assert any("text at 10 pt" in note for note in notes(run_audit(deck, tmp_path)))
