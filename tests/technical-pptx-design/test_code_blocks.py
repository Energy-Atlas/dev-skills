"""Tests for skills/technical-pptx-design/scripts/code_blocks.py: code listings and comparisons as native tables.

Run from the repository root: python -m pytest tests (needs python-pptx, Pygments, pytest).
"""
import sys
from pathlib import Path

import pytest
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).parents[2] / "skills" / "technical-pptx-design" / "scripts"))
from audit_pptx import audit, problem_count  # noqa: E402
from code_blocks import (DARK, LIGHT, NO_STYLE_NO_GRID, add_code_block, add_code_diff, diff_hunks,  # noqa: E402
                         highlighted_lines, monospace_font)
from pptx_helpers import new_widescreen_deck, remove_unused_layouts, set_theme_font  # noqa: E402

OLD = "def area(w, h):\n    return w * h\n"
NEW = "def area(w, h):\n    \"\"\"Floor area in m2.\"\"\"\n    return w * h\n"


def slide_of(deck):
    slide = deck.slides.add_slide(next(layout for layout in deck.slide_layouts if layout.name == "Title Only"))
    slide.shapes.title.text = "Code"
    return slide


def cell_text(table, row, column):
    return table.cell(row, column).text_frame.text


def fill(table, row, column):
    return str(table.cell(row, column).fill.fore_color.rgb)


def ratio(a, b):
    def luminance(colour):
        channels = [int(colour[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        channels = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in channels]
        return 0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2]
    high, low = sorted((luminance(a), luminance(b)), reverse=True)
    return (high + 0.05) / (low + 0.05)


@pytest.mark.parametrize("palette", [LIGHT, DARK])
def test_every_code_colour_reads_on_every_backdrop(palette):
    texts = ["text", "muted", "comment", "keyword", "string", "number", "function", "added", "removed"]
    for backdrop in ("backdrop", "label", "green", "red"):
        for text in texts:
            assert ratio(palette[text], palette[backdrop]) >= 4.5, (text, backdrop)


def test_monospace_font_takes_the_first_installed_family():
    assert monospace_font(installed={"Consolas", "Courier New"}) == "Consolas"
    assert monospace_font(installed={"Cascadia Mono Regular", "Consolas"}) == "Cascadia Mono"
    assert monospace_font(installed=set()) == "Courier New"


def test_highlighting_keeps_multiline_strings_and_expands_tabs():
    lines = highlighted_lines('x = """a\nb"""\n\tif x:\n', "python")
    assert len(lines) == 3
    assert {category for category, _ in lines[1]} == {"string"}
    assert "".join(text for _, text in lines[1]) == 'b"""'
    assert "".join(text for _, text in lines[2]).startswith("    if")


def test_diff_hunks_pair_changes():
    assert diff_hunks(["a", "b", "c"], ["a", "x", "c"]) == [("equal", [(0, 0)]), ("change", [1], [1]),
                                                             ("equal", [(2, 2)])]


def test_code_block_is_a_native_table_with_numbers_marks_and_the_code_font():
    deck = new_widescreen_deck()
    shape = add_code_block(slide_of(deck), NEW, "python", Inches(0.6), Inches(1.6), Inches(8), first_line=40,
                           marks={41: "green"}, label="geometry.py", font="Consolas")
    table = shape.table
    assert shape._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId")).text == NO_STYLE_NO_GRID
    assert len(table.rows) == 4  # label and three lines
    assert cell_text(table, 0, 0) == "geometry.py"
    assert [cell_text(table, row, 0) for row in (1, 2, 3)] == ["40", "41", "42"]
    assert cell_text(table, 3, 1) == "    return w * h"
    assert fill(table, 2, 0) == fill(table, 2, 1) == LIGHT["green"]  # the whole line
    assert fill(table, 1, 1) == LIGHT["backdrop"]
    runs = [run for row in table.rows for cell in row.cells for p in cell.text_frame.paragraphs for run in p.runs]
    assert {run.font.name for run in runs} == {"Consolas"}
    assert min(run.font.size.pt for run in runs) >= 12


def test_unified_diff_marks_removed_and_added_lines():
    deck = new_widescreen_deck()
    table = add_code_diff(slide_of(deck), OLD, NEW, "python", Inches(0.6), Inches(1.6), Inches(12),
                          labels=("v1", "v2"), font="Consolas").table
    assert cell_text(table, 0, 0) == "v1 → v2"
    rows = [(cell_text(table, r, 0).strip(), cell_text(table, r, 1).strip(), cell_text(table, r, 2))
            for r in range(1, len(table.rows))]
    assert rows == [("1", "1", '  def area(w, h):'), ("", "2", '+     """Floor area in m2."""'),
                    ("2", "3", "      return w * h")]
    assert fill(table, 2, 2) == LIGHT["green"]


def test_split_diff_lines_up_both_sides_at_one_size():
    deck = new_widescreen_deck()
    old_shape, new_shape = add_code_diff(slide_of(deck), OLD, NEW, "python", Inches(0.6), Inches(1.6), Inches(12),
                                         layout="split", font="Consolas", dark=True)
    old, new = old_shape.table, new_shape.table
    assert len(old.rows) == len(new.rows) == 3
    assert fill(old, 1, 1) == DARK["label"] and cell_text(old, 1, 1).strip() == ""  # stands in for the added line
    assert cell_text(new, 1, 1).startswith("+")
    assert cell_text(old, 2, 1)[2:] == cell_text(new, 2, 1)[2:] == "    return w * h"
    size = lambda t: t.cell(0, 1).text_frame.paragraphs[0].runs[0].font.size
    assert size(old) == size(new)
    assert old_shape.left + old_shape.width < new_shape.left


def test_long_lines_shrink_then_warn():
    deck = new_widescreen_deck()
    line = "x = " + "a + " * 30 + "a\n"
    with pytest.warns(UserWarning, match="wrap"):
        shape = add_code_block(slide_of(deck), line, "python", Inches(0.6), Inches(1.6), Inches(4), font="Consolas")
    assert shape.table.cell(0, 1).text_frame.paragraphs[0].runs[0].font.size == Pt(12)


def test_audit_accepts_code_blocks_with_the_declared_code_font(tmp_path):
    deck = new_widescreen_deck()
    add_code_block(slide_of(deck), NEW, "python", Inches(0.6), Inches(1.6), Inches(8), font="Consolas")
    set_theme_font(deck, "Calibri")
    remove_unused_layouts(deck)
    deck.save(tmp_path / "deck.pptx")
    assert problem_count(audit(tmp_path / "deck.pptx", code_font="Consolas")) == 0
    assert problem_count(audit(tmp_path / "deck.pptx")) > 0  # undeclared, the code font is a stray font


def test_context_keeps_only_the_lines_around_a_change():
    deck = new_widescreen_deck()
    old = "".join(f"x{n} = {n}\n" for n in range(1, 21))
    new = old.replace("x10 = 10\n", "x10 = 100\n")
    table = add_code_diff(slide_of(deck), old, new, "python", Inches(0.6), Inches(1.6), Inches(12), context=2,
                          font="Consolas").table
    code = [cell_text(table, r, 2) for r in range(len(table.rows))]
    assert code == ["⋯", "  x8 = 8", "  x9 = 9", "- x10 = 10", "+ x10 = 100", "  x11 = 11", "  x12 = 12", "⋯"]
    assert fill(table, 0, 2) == LIGHT["label"]
